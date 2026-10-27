"""The downloadable underwriting template (GET /deals/import/template).

One workbook: Instructions, Property, Unit Mix and Financing tabs, written
from the field catalog so the template and the importer can't drift apart.
Inputs are shaded, numbers carry $ / % formats, and the choice fields get
dropdowns fed from a hidden Lists sheet. Sample values describe a typical
suburban garden deal for the user to overwrite.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.reports import theme
from app.spreadsheet.fields import (
    FIELDS,
    SAMPLE_UNIT_MIX,
    UNIT_MIX_COLUMNS,
    FieldSpec,
    Section,
)

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILENAME = "dealrank_underwriting_template.xlsx"

FONT = "Calibri"
_thin = Side(style="thin", color=theme.RULE)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEADER_FILL = PatternFill("solid", fgColor=theme.PRIMARY)
BAND_FILL = PatternFill("solid", fgColor=theme.TINT)
INPUT_FILL = PatternFill("solid", fgColor=theme.HIGHLIGHT)

MONEY_FORMAT = '"$"#,##0'
RENT_FORMAT = '"$"#,##0.00'
PERCENT_FORMAT = "0.00%"
INTEGER_FORMAT = "#,##0"

_NUMBER_FORMATS = {
    "money": MONEY_FORMAT,
    "percent": PERCENT_FORMAT,
    "integer": INTEGER_FORMAT,
    "number": "0",
    "text": "@",
    "choice": "@",
}


def _title(ws: Worksheet, text: str, subtitle: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=theme.INK)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color=theme.MUTED)
    ws.row_dimensions[1].height = 26


def _fit_to_page_width(ws: Worksheet) -> None:
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = Font(name=FONT, bold=True, color=theme.WHITE)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20


def _band(ws: Worksheet, row: int, text: str, width: int) -> None:
    for col in range(1, width + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = BAND_FILL
        cell.border = BORDER
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, bold=True, color=theme.PRIMARY_MID)


def _sample_value(spec: FieldSpec) -> object:
    if spec.sample == "":
        return None
    if spec.kind == "choice":
        return next(c.label for c in spec.choices if c.value == spec.sample)
    if spec.kind == "text":
        return spec.sample
    number = Decimal(spec.sample)
    if spec.kind == "percent":
        return float(number / 100)
    if spec.kind == "integer":
        return int(number)
    return float(number) if number != number.to_integral_value() else int(number)


def _lists_sheet(wb: Workbook) -> dict[str, str]:
    """Writes each choice field's labels to a hidden sheet; returns the
    dropdown range for each field key."""
    ws = wb.create_sheet("Lists")
    ranges: dict[str, str] = {}
    choice_fields = [f for f in FIELDS if f.kind == "choice"]
    for col, spec in enumerate(choice_fields, start=1):
        ws.cell(row=1, column=col, value=spec.key)
        for row, choice in enumerate(spec.choices, start=2):
            ws.cell(row=row, column=col, value=choice.label)
        letter = get_column_letter(col)
        ranges[spec.key] = f"Lists!${letter}$2:${letter}${len(spec.choices) + 1}"
    ws.sheet_state = "hidden"
    return ranges


def _fields_sheet(
    wb: Workbook,
    title: str,
    heading: str,
    subtitle: str,
    sections: tuple[Section, ...],
    dropdowns: dict[str, str],
) -> None:
    ws = wb.create_sheet(title)
    _fit_to_page_width(ws)
    _title(ws, heading, subtitle)
    _header(ws, 4, ["Field", "Value", "Unit", "Notes"])
    for letter, width in zip("ABCD", (34, 22, 16, 58), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A5"

    row = 5
    group = None
    for spec in (f for f in FIELDS if f.section in sections):
        if spec.group != group:
            if group is not None:
                row += 1
            group = spec.group
            _band(ws, row, group, 4)
            row += 1
        label = ws.cell(row=row, column=1, value=spec.label)
        label.font = Font(name=FONT, color=theme.INK)
        label.border = BORDER

        value = ws.cell(row=row, column=2, value=_sample_value(spec))
        value.fill = INPUT_FILL
        value.border = BORDER
        value.font = Font(name=FONT, bold=True, color=theme.INK)
        value.number_format = (
            RENT_FORMAT
            if spec.kind == "money" and "month" in spec.unit
            else _NUMBER_FORMATS[spec.kind]
        )
        value.alignment = Alignment(
            horizontal="left" if spec.kind in ("text", "choice") else "right"
        )
        if spec.key in dropdowns:
            dv = DataValidation(
                type="list",
                formula1=f"={dropdowns[spec.key]}",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Pick from the list",
                error="Choose one of the options in the dropdown.",
            )
            ws.add_data_validation(dv)
            dv.add(value)

        unit = ws.cell(row=row, column=3, value=spec.unit or None)
        unit.font = Font(name=FONT, color=theme.MUTED)
        unit.alignment = Alignment(indent=1)
        unit.border = BORDER
        notes = ws.cell(row=row, column=4, value=spec.notes or None)
        notes.font = Font(name=FONT, italic=True, color=theme.MUTED)
        notes.alignment = Alignment(indent=1)
        notes.border = BORDER
        row += 1


def _unit_mix_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Unit Mix")
    _fit_to_page_width(ws)
    _title(
        ws,
        "Unit mix",
        "One row per unit type. Add rows above the total; leave unused rows blank.",
    )
    headers = [c.label for c in UNIT_MIX_COLUMNS]
    _header(ws, 4, headers)
    for letter, width in zip("ABCDE", (28, 12, 16, 14, 16), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A5"

    formats = ["@", INTEGER_FORMAT, INTEGER_FORMAT, INTEGER_FORMAT, MONEY_FORMAT]
    first, rows = 5, 10
    for i in range(rows):
        sample = SAMPLE_UNIT_MIX[i] if i < len(SAMPLE_UNIT_MIX) else (None,) * 5
        for col, (value, fmt) in enumerate(zip(sample, formats, strict=True), start=1):
            cell = ws.cell(row=first + i, column=col, value=value)
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.number_format = fmt
            cell.font = Font(name=FONT, color=theme.INK, bold=col > 1)
    last = first + rows - 1
    total_row = last + 1
    totals = {
        1: "Total",
        2: f"=SUM(B{first}:B{last})",
        4: f"=IFERROR(SUMPRODUCT(B{first}:B{last},D{first}:D{last})/B{total_row},0)",
        5: f"=IFERROR(SUMPRODUCT(B{first}:B{last},E{first}:E{last})/B{total_row},0)",
    }
    for col in range(1, 6):
        cell = ws.cell(row=total_row, column=col, value=totals.get(col))
        cell.font = Font(name=FONT, bold=True, color=theme.INK)
        cell.fill = BAND_FILL
        cell.border = BORDER
        cell.number_format = formats[col - 1]
    ws.cell(
        row=total_row + 2,
        column=1,
        value="Units and average rent on the Property tab are filled from this table.",
    ).font = Font(name=FONT, italic=True, color=theme.MUTED)


def _instructions_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instructions"
    _fit_to_page_width(ws)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 96
    ws.sheet_view.showGridLines = False
    ws["B2"] = f"{theme.NAME} underwriting template"
    ws["B2"].font = Font(name=FONT, size=18, bold=True, color=theme.INK)
    ws["B3"] = "Fill in a deal here, then import it from the New deal form."
    ws["B3"].font = Font(name=FONT, size=11, color=theme.MUTED)

    steps = [
        "Replace the sample values in the shaded cells on the Property, Unit Mix and "
        "Financing tabs.",
        "Leave a cell blank to keep the form's default. Optional sections can stay at 0.",
        "Enter percentages as percentages (5.75%). Dropdown cells only take a listed option.",
        "In DealRank, open New deal, choose Import spreadsheet and pick this file.",
        "Check the values on the review screen, then fill the form and save the deal.",
    ]
    ws["B5"] = "How to use it"
    ws["B5"].font = Font(name=FONT, size=12, bold=True, color=theme.PRIMARY_MID)
    for i, step in enumerate(steps, start=1):
        cell = ws.cell(row=5 + i, column=2, value=f"{i}.  {step}")
        cell.font = Font(name=FONT, color=theme.INK)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    row = 6 + len(steps) + 1
    ws.cell(row=row, column=1).fill = INPUT_FILL
    ws.cell(row=row, column=1).border = BORDER
    ws.cell(row=row, column=2, value="Shaded cells are inputs.").font = Font(
        name=FONT, color=theme.INK
    )
    ws.cell(
        row=row + 2,
        column=2,
        value=(
            "Have your own underwriting model? You can import that too: DealRank looks for "
            "familiar labels (Purchase Price, Vacancy, Exit Cap Rate, a unit mix table) and "
            "lets you confirm or fix each value before anything is filled in."
        ),
    ).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row + 2].height = 32


def build_template_bytes() -> bytes:
    wb = Workbook()
    _instructions_sheet(wb)
    dropdowns = _lists_sheet(wb)
    _fields_sheet(
        wb,
        "Property",
        "Property and underwriting",
        "Year 1 assumptions for the deal.",
        ("property", "ancillary"),
        dropdowns,
    )
    _unit_mix_sheet(wb)
    _fields_sheet(
        wb,
        "Financing",
        "Financing",
        "The senior loan, plus any mezzanine debt or preferred equity.",
        ("financing",),
        dropdowns,
    )
    # Lists was created second for the dropdown ranges; keep it last.
    wb.move_sheet("Lists", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Lists"))
    wb.properties.title = f"{theme.NAME} underwriting template"
    wb.properties.creator = theme.NAME

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
