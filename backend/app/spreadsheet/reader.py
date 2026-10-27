"""Loads an uploaded workbook into plain cells, and parses cell values.

Supports .xlsx/.xlsm (openpyxl, read-only, cached formula values) and
.csv/.tsv. Only the top-left of each sheet is scanned (MAX_ROWS x MAX_COLS)
so a huge export can't stall a request.
"""

from __future__ import annotations

import csv
import datetime
import io
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_BYTES = 5 * 1024 * 1024
MAX_SHEETS = 12
MAX_ROWS = 400
MAX_COLS = 40

EXCEL_SUFFIXES = (".xlsx", ".xlsm")
TEXT_SUFFIXES = (".csv", ".tsv", ".txt")


class SpreadsheetError(ValueError):
    """The file can't be read; the message is shown to the user."""


@dataclass(frozen=True)
class Cell:
    sheet: str
    row: int  # 1-based, as Excel shows it
    col: int  # 1-based
    value: object  # str, int, float, Decimal, bool or datetime
    percent_format: bool = False

    @property
    def ref(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def address(self) -> str:
        return f"{self.sheet}!{self.ref}"

    @property
    def text(self) -> str:
        if isinstance(self.value, str):
            return self.value.strip()
        if isinstance(self.value, datetime.date | datetime.datetime):
            return self.value.isoformat()[:10]
        number = parse_number(self.value)
        if number is None:
            return str(self.value)
        if self.percent_format:
            return f"{format_decimal(number * 100)}%"
        return format_decimal(number, thousands=True)


@dataclass
class Sheet:
    name: str
    cells: dict[tuple[int, int], Cell]

    def get(self, row: int, col: int) -> Cell | None:
        return self.cells.get((row, col))


def load_sheets(filename: str, data: bytes) -> list[Sheet]:
    if len(data) > MAX_BYTES:
        raise SpreadsheetError("That file is over 5 MB. Upload a smaller workbook.")
    if not data:
        raise SpreadsheetError("That file is empty.")
    lower = filename.lower()
    if lower.endswith(".xls"):
        raise SpreadsheetError(
            "Old .xls files aren't supported. Open it in Excel and save it as .xlsx."
        )
    if lower.endswith(EXCEL_SUFFIXES):
        return _load_workbook(data)
    if lower.endswith(TEXT_SUFFIXES):
        return [_load_delimited(filename, data)]
    raise SpreadsheetError("Upload an Excel workbook (.xlsx) or a .csv file.")


def _load_workbook(data: bytes) -> list[Sheet]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of zip/XML errors
        raise SpreadsheetError("That file couldn't be opened as an Excel workbook.") from exc
    sheets: list[Sheet] = []
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            cells: dict[tuple[int, int], Cell] = {}
            for row in worksheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS):
                for cell in row:
                    value = getattr(cell, "value", None)
                    if value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    fmt = getattr(cell, "number_format", "") or ""
                    cells[(cell.row, cell.column)] = Cell(
                        worksheet.title, cell.row, cell.column, value, "%" in fmt
                    )
            sheets.append(Sheet(worksheet.title, cells))
    finally:
        workbook.close()
    return sheets


def _load_delimited(filename: str, data: bytes) -> Sheet:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    first_line = text.split("\n", 1)[0]
    delimiter = "\t" if filename.lower().endswith(".tsv") or "\t" in first_line else ","
    name = re.sub(r"\.[^.]+$", "", filename.rsplit("/", 1)[-1]) or "Sheet1"
    cells: dict[tuple[int, int], Cell] = {}
    for r, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        if r > MAX_ROWS:
            break
        for c, value in enumerate(row[:MAX_COLS], start=1):
            if value.strip():
                cells[(r, c)] = Cell(name, r, c, value.strip())
    return Sheet(name, cells)


_SUFFIXES = {"k": Decimal(1_000), "m": Decimal(1_000_000), "mm": Decimal(1_000_000)}
_NUMBER_RE = re.compile(r"^(-?\d+(?:\.\d+)?|-?\.\d+)(k|m|mm)?$")


def parse_number(value: object) -> Decimal | None:
    """A cell value as a Decimal: handles $, commas, spaces, a trailing %,
    (1,234) negatives and k / M / MM suffixes. None when it isn't a number."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float | Decimal):
        try:
            number = Decimal(str(value))
        except InvalidOperation:
            return None
        return number if number.is_finite() else None
    if not isinstance(value, str):
        return None
    text = value.strip().lower().replace(",", "").replace("$", "").replace(" ", "")
    text = text.rstrip("%")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    match = _NUMBER_RE.match(text)
    if not match:
        return None
    number = Decimal(match.group(1))
    if match.group(2):
        number *= _SUFFIXES[match.group(2)]
    return -number if negative else number


def has_percent_sign(value: object) -> bool:
    return isinstance(value, str) and value.strip().endswith("%")


def format_decimal(number: Decimal, thousands: bool = False) -> str:
    """Plain decimal text with no exponent or trailing zeros, e.g. 5.75 or 24500000."""
    rounded = round(number, 6)
    text = f"{rounded:,f}" if thousands else f"{rounded:f}"
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return "0" if text in ("-0", "") else text
