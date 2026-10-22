"""PPTX export. See docs/build-plan.md 6.3 — same content as the PDF, as
slides. This is the artifact that gets presented, so it carries the
methodology note about unlevered-IRR ranking directly on the ranking slide,
not just in an appendix.

Built as a 16:9 investment-committee deck: cover, executive summary (KPI
tiles, findings, return-vs-hurdle chart), ranking, the comparison split into
operating and returns slides, how each hurdle was built, then per deal a
profile, capital stack and pro forma, closing with methodology. Every slide
carries a one-sentence action title so the deck reads from the titles alone.
"""

import io
from decimal import Decimal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.api.ranking_data import RankedEntry
from app.engine.types import LayerGroup
from app.reports import theme
from app.reports.assumptions import ASSUMPTION_GROUPS, format_assumption_value, present_rows
from app.reports.capital_stack import (
    CAPITAL_STACK_HEADERS,
    capital_stack_rows,
    capital_stack_segments,
)
from app.reports.data import ReportData
from app.reports.formatting import money, multiple, number, percent, short_name


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


INK, PRIMARY, PRIMARY_MID, ACCENT = (
    _rgb(theme.INK),
    _rgb(theme.PRIMARY),
    _rgb(theme.PRIMARY_MID),
    _rgb(theme.ACCENT),
)
TINT, STRIPE, RULE, MUTED, WHITE = (
    _rgb(theme.TINT),
    _rgb(theme.STRIPE),
    _rgb(theme.RULE),
    _rgb(theme.MUTED),
    _rgb(theme.WHITE),
)

FONT = "Calibri"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
CONTENT_W = SLIDE_W - 2 * MARGIN
CONTENT_TOP = 1.7
_BLANK_LAYOUT_INDEX = 6  # "Blank" in the default python-pptx template
_NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

_GROUP_COLORS = {
    LayerGroup.SENIOR: PRIMARY,
    LayerGroup.SUBORDINATE: PRIMARY_MID,
    LayerGroup.EQUITY: ACCENT,
}


# ---------------------------------------------------------------- primitives


def _rect(slide, x, y, w, h, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    # Drop the theme style reference so no inherited shadow or outline is drawn.
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def _text(
    slide,
    x,
    y,
    w,
    h,
    text: str,
    size: float = 12,
    bold: bool = False,
    color: RGBColor = INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    italic: bool = False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _font(run, size, bold, color, italic)
    return box


def _font(run, size: float, bold: bool, color: RGBColor, italic: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _paragraphs(slide, x, y, w, h, lines: list[str], size=11, color=INK, italic=False, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.space_after = Pt(gap)
        run = paragraph.add_run()
        run.text = line
        _font(run, size, False, color, italic)
    return box


class _Deck:
    """The presentation plus the page counter shown in each footer."""

    def __init__(self, data: ReportData) -> None:
        self.data = data
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.page = 0

    def blank(self):
        self.page += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[_BLANK_LAYOUT_INDEX])

    def content_slide(self, kicker: str, title: str):
        slide = self.blank()
        _rect(slide, MARGIN, 0.42, 0.12, 0.12, ACCENT)
        _text(slide, MARGIN + 0.22, 0.38, 6, 0.25, kicker.upper(), 10, True, PRIMARY_MID)
        _text(slide, MARGIN, 0.7, CONTENT_W, 0.85, title, 24, True, PRIMARY)
        line = _rect(slide, MARGIN, 7.0, CONTENT_W, 0.01, RULE)
        line.name = "footer rule"
        _text(slide, MARGIN, 7.08, 1.5, 0.25, theme.NAME.upper(), 9, True, PRIMARY)
        _text(
            slide,
            MARGIN + 1.5,
            7.08,
            8,
            0.25,
            f"{self.data.generated_on}  |  Screening analysis, not an appraisal",
            9,
            False,
            MUTED,
        )
        _text(
            slide,
            SLIDE_W - MARGIN - 1,
            7.08,
            1,
            0.25,
            str(self.page),
            9,
            True,
            PRIMARY,
            align=PP_ALIGN.RIGHT,
        )
        return slide


# -------------------------------------------------------------------- tables


def _bottom_border(cell, color: str, width_pt: float = 0.5) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("a:lnB")):
        tc_pr.remove(old)
    line = parse_xml(
        f'<a:lnB xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'w="{int(width_pt * 12700)}" cap="flat" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:prstDash val="solid"/></a:lnB>'
    )
    # Line properties precede the fill in the cell-properties schema.
    tc_pr.insert(0, line)


def _style_cell(
    cell,
    text: str,
    size: float,
    bold: bool = False,
    color: RGBColor = INK,
    fill: RGBColor | None = None,
    align=PP_ALIGN.LEFT,
    border: str | None = theme.RULE,
) -> None:
    cell.text = text
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            _font(run, size, bold, color)
    cell.text_frame.word_wrap = True
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    if border:
        _bottom_border(cell, border)


def _table(
    slide,
    x: float,
    y: float,
    widths: list[float],
    rows: list[list[str]],
    row_h: float,
    size: float,
    aligns: list | None = None,
    section_rows: frozenset[int] = frozenset(),
    bold_rows: frozenset[int] = frozenset(),
    highlight_rows: frozenset[int] = frozenset(),
    header_h: float | None = None,
    bold_cols: frozenset[int] = frozenset(),
):
    """A borderless table with a dark header band and hairline row rules.
    Row 0 is the header. Section rows span every column as a tinted label;
    highlight rows carry the ranking-basis treatment. Geometry is stored
    exactly (widths and row heights) so it can be checked against the slide."""
    n_rows, n_cols = len(rows), len(widths)
    aligns = aligns or [PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * (n_cols - 1)
    header_h = header_h or row_h
    shape = slide.shapes.add_table(
        n_rows,
        n_cols,
        Inches(x),
        Inches(y),
        Inches(sum(widths)),
        Inches(header_h + row_h * (n_rows - 1)),
    )
    table = shape.table
    table.horz_banding = False
    style_id = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style_id is not None:
        style_id.text = _NO_STYLE_NO_GRID
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    for row in range(n_rows):
        table.rows[row].height = Inches(header_h if row == 0 else row_h)

    for r, cells in enumerate(rows):
        if r == 0:
            for c, value in enumerate(cells):
                _style_cell(
                    table.cell(0, c), value, size, True, WHITE, PRIMARY, aligns[c], border=None
                )
        elif r in section_rows:
            table.cell(r, 0).merge(table.cell(r, n_cols - 1))
            for c in range(n_cols):
                _style_cell(
                    table.cell(r, c),
                    cells[0] if c == 0 else "",
                    size - 1.5,
                    True,
                    PRIMARY,
                    TINT,
                    PP_ALIGN.LEFT,
                    border=None,
                )
        else:
            emphasised = r in bold_rows
            highlighted = r in highlight_rows
            for c, value in enumerate(cells):
                _style_cell(
                    table.cell(r, c),
                    value,
                    size,
                    emphasised or highlighted or c in bold_cols,
                    PRIMARY if highlighted else INK,
                    TINT if highlighted else (STRIPE if emphasised else None),
                    aligns[c],
                    border=theme.PRIMARY_MID if emphasised else theme.RULE,
                )
    return shape


# ------------------------------------------------------------------- widgets


def _kpi_tiles(slide, y: float, items: list[tuple[str, str]], height: float = 1.0) -> None:
    gap = 0.12
    width = (CONTENT_W - gap * (len(items) - 1)) / len(items)
    for i, (label, value) in enumerate(items):
        x = MARGIN + i * (width + gap)
        _rect(slide, x, y, width, height, STRIPE)
        _rect(slide, x, y, width, 0.05, ACCENT)
        _text(
            slide,
            x,
            y + 0.17,
            width,
            0.45,
            value,
            24,
            True,
            PRIMARY,
            PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE,
        )
        _text(slide, x, y + 0.66, width, 0.25, label, 10, False, MUTED, PP_ALIGN.CENTER)


def _kpis_for(entry: RankedEntry) -> list[tuple[str, str]]:
    m = entry.metrics
    return [
        ("Year 1 NOI", money(m.year_one_noi)),
        ("Going-in cap rate", percent(m.going_in_cap_rate)),
        ("Unlevered IRR", percent(m.unlevered_irr)),
        ("Levered IRR (reference)", percent(m.levered_irr)),
        ("Equity multiple", multiple(m.equity_multiple)),
        ("Lowest DSCR", number(m.min_dscr)),
    ]


def _style_chart(chart, size: float = 10) -> None:
    chart.font.size = Pt(size)
    chart.font.name = FONT
    chart.font.color.rgb = INK


def _return_chart(slide, data: ReportData, x, y, w, h) -> None:
    perf = data.perf
    chart_data = CategoryChartData()
    chart_data.categories = [p.name for p in perf]
    chart_data.add_series("Unlevered IRR", [float(p.irr) for p in perf])
    chart_data.add_series("Hurdle", [float(p.hurdle) for p in perf])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
    ).chart
    _style_chart(chart)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.overlap = -5
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "0.0%"
    labels.number_format_is_linked = False
    labels.font.size = Pt(10)
    for series, color in zip(plot.series, (PRIMARY_MID, RULE), strict=True):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.invert_if_negative = False
    # First deal on top, and no value axis: the data labels carry the numbers.
    chart.category_axis.reverse_order = True
    chart.category_axis.format.line.color.rgb = RULE
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False


def _flow_chart(slide, entry: RankedEntry, x, y, w, h) -> None:
    flows = entry.metrics.annual_cash_flows[:_MAX_PROFORMA_YEARS]
    chart_data = CategoryChartData()
    chart_data.categories = [f"Yr {f.year}" for f in flows]
    chart_data.add_series("Net operating income", [float(f.noi) for f in flows])
    chart_data.add_series(
        "Cash flow after debt service", [float(f.levered_cash_flow) for f in flows]
    )
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
    ).chart
    _style_chart(chart, 9)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.gap_width = 70
    for series, color in zip(plot.series, (PRIMARY, ACCENT), strict=True):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.invert_if_negative = False
    chart.value_axis.tick_labels.number_format = '$#,##0,"k"'
    chart.value_axis.tick_labels.number_format_is_linked = False
    chart.value_axis.major_gridlines.format.line.color.rgb = RULE
    chart.value_axis.format.line.fill.background()
    chart.category_axis.format.line.color.rgb = RULE
    # Keep year labels clear of bars when cash flow after debt service is negative.
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW


# -------------------------------------------------------------------- slides


def _year_one(entry: RankedEntry):
    return entry.metrics.annual_cash_flows[0]


def _standing(entry: RankedEntry) -> str:
    return f"Rank {entry.rank}" if entry.rank else "Excluded"


def _add_cover(deck: _Deck) -> None:
    data = deck.data
    slide = deck.blank()
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, PRIMARY)
    _rect(slide, 0, 0, 0.18, SLIDE_H, ACCENT)
    _rect(slide, 0.9, 0.78, 0.14, 0.14, ACCENT)
    _text(slide, 1.15, 0.72, 4, 0.3, theme.NAME.upper(), 13, True, WHITE)
    _text(slide, 0.9, 2.0, 9, 0.35, "INVESTMENT ANALYSIS", 13, True, ACCENT)
    _text(slide, 0.9, 2.4, 11, 1.7, "Deal Comparison\n& Ranking", 48, True, WHITE)
    _text(slide, 0.9, 4.25, 11, 0.4, data.subtitle, 15, False, _rgb("C9DAD5"))
    _text(slide, 0.9, 4.7, 11, 0.35, data.generated_on, 13, False, _rgb("C9DAD5"))

    # Properties analysed, as a light-on-dark table.
    rows = [["", "Property", "Class", "Standing"]]
    for i, e in enumerate(data.entries, start=1):
        rows.append(
            [
                f"{i:02d}",
                e.deal.name,
                format_assumption_value("sub_asset_class", e.deal.sub_asset_class),
                _standing(e),
            ]
        )
    row_h = min(0.3, 1.45 / len(rows))
    shape = slide.shapes.add_table(
        len(rows),
        4,
        Inches(0.9),
        Inches(5.35),
        Inches(8.6),
        Inches(row_h * len(rows)),
    )
    table = shape.table
    table.horz_banding = False
    style_id = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style_id is not None:
        style_id.text = _NO_STYLE_NO_GRID
    for col, width in enumerate((0.6, 4.4, 2.3, 1.3)):
        table.columns[col].width = Inches(width)
    for r, cells in enumerate(rows):
        table.rows[r].height = Inches(row_h)
        for c, value in enumerate(cells):
            _style_cell(
                table.cell(r, c),
                value,
                10 if r else 9,
                bold=(r == 0 or c in (0, 3)),
                color=ACCENT if (r == 0 or c == 0) else WHITE,
                align=PP_ALIGN.RIGHT if c == 3 else PP_ALIGN.LEFT,
                border="2F6B60",
            )


def _add_summary_slide(deck: _Deck) -> None:
    data = deck.data
    leader = next((e for e in data.entries if e.rank), None)
    if data.perf:
        top = data.perf[0]
        title = (
            f"{top.name} leads at {percent(top.irr)} unlevered IRR, "
            f"{abs(top.spread) * 10000:,.0f} bps "
            f"{'above' if top.spread >= 0 else 'below'} its hurdle"
        )
    else:
        title = "No deal cleared the screening gates"
    slide = deck.content_slide("Executive summary", title)

    top_of_body = CONTENT_TOP
    if leader is not None:
        _kpi_tiles(slide, top_of_body, _kpis_for(leader))
        _text(
            slide,
            MARGIN,
            top_of_body + 1.07,
            8,
            0.25,
            f"Top-ranked deal: {leader.deal.name}",
            10,
            False,
            MUTED,
            italic=True,
        )
        top_of_body += 1.45

    body_h = 6.85 - top_of_body
    _text(slide, MARGIN, top_of_body + 0.05, 5.4, 0.3, "Key findings", 14, True, PRIMARY)
    _rect(slide, MARGIN, top_of_body + 0.42, 0.45, 0.04, ACCENT)
    _paragraphs(
        slide,
        MARGIN,
        top_of_body + 0.6,
        5.4,
        body_h - 0.6,
        [f"•  {t}" for t in data.takeaways],
        size=13,
        gap=9,
    )
    if data.perf:
        _text(
            slide,
            6.4,
            top_of_body + 0.05,
            6.4,
            0.3,
            "Return vs. hurdle" + (" (before costs)" if data.mode == "risk" else ""),
            14,
            True,
            PRIMARY,
        )
        _return_chart(slide, data, 6.3, top_of_body + 0.4, 6.55, body_h - 0.4)


def _add_ranking_slide(deck: _Deck) -> None:
    data = deck.data
    slide = deck.content_slide("Ranking", "Ranked on unlevered IRR against each deal's hurdle")
    rows = [data.ranking_headers] + data.ranking_rows
    cols = len(data.ranking_headers)
    first, last = 0.8, 3.6
    rest = (CONTENT_W - first - last) / (cols - 2)
    row_h = min(0.52, 3.6 / len(rows))
    shape = _table(
        slide,
        MARGIN,
        CONTENT_TOP,
        [first, last] + [rest] * (cols - 2),
        rows,
        row_h,
        13 if cols <= 6 else 12,
        aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * (cols - 2),
        highlight_rows=frozenset({1}) if data.ranking_rows else frozenset(),
        bold_cols=frozenset({0, 2}),
        header_h=0.6,
    )
    bottom = CONTENT_TOP + 0.6 + row_h * (len(rows) - 1)

    lines = [data.basis_note]
    lines += [f"Excluded: {name} — {reason}" for name, reason in data.excluded]
    lines += [f"Note: {name} — {warning}" for name, warning in data.warnings]
    _paragraphs(
        slide, MARGIN, bottom + 0.3, CONTENT_W, 6.9 - bottom - 0.3, lines, size=10.5, color=MUTED
    )
    del shape


def _add_operating_slide(deck: _Deck) -> None:
    data = deck.data
    entries = data.entries
    slide = deck.content_slide(
        "Comparison", "Year 1 operating waterfall: from gross rent to net operating income"
    )
    lines = [
        ("Hold period", lambda e: f"{e.deal.hold_period_years} yrs", False),
        ("Gross potential rent", lambda e: money(_year_one(e).gpr), False),
        ("Less: vacancy loss", lambda e: money(_year_one(e).vacancy_loss), False),
        ("Less: bad debt & concessions", lambda e: money(_year_one(e).credit_loss), False),
        ("Effective gross income", lambda e: money(_year_one(e).egi), True),
        ("Less: operating expenses", lambda e: money(_year_one(e).opex), False),
        ("Less: turnover expense", lambda e: money(_year_one(e).turnover_expense), False),
        ("Less: management fee", lambda e: money(_year_one(e).management_fee), False),
        ("Year 1 NOI", lambda e: money(e.metrics.year_one_noi), True),
        ("Less: replacement reserves", lambda e: money(_year_one(e).replacement_reserves), False),
        ("Debt service (all debt)", lambda e: money(_year_one(e).debt_service), False),
        ("Preferred return", lambda e: money(_year_one(e).preferred_distribution), False),
    ]
    _comparison_table(slide, entries, lines, size=11)


def _add_returns_slide(deck: _Deck) -> None:
    data = deck.data
    entries = data.entries
    risk = data.mode == "risk"
    slide = deck.content_slide(
        "Comparison", "Returns, credit and exit: what each deal earns and how it is financed"
    )
    # (label, extractor, bold) — a (label, None, False) entry is a section row.
    lines: list = [("Returns", None, False)]
    lines.append(
        (
            "Unlevered IRR (after costs)" if risk else data.basis_label,
            lambda e: percent(e.metrics.unlevered_irr),
            False,
        )
    )
    if risk:
        lines.append(
            (data.basis_label, lambda e: percent(e.metrics.unlevered_irr_before_costs), False)
        )
    lines += [
        ("Levered IRR (reference)", lambda e: percent(e.metrics.levered_irr), False),
        ("Going-in cap rate", lambda e: percent(e.metrics.going_in_cap_rate), False),
        ("Equity invested", lambda e: money(e.metrics.equity_invested), False),
        ("Cash-on-cash", lambda e: percent(e.metrics.cash_on_cash), False),
        ("Equity multiple", lambda e: multiple(e.metrics.equity_multiple), True),
        ("Credit", None, False),
        ("Year 1 senior DSCR", lambda e: number(e.metrics.year_one_senior_dscr), False),
        ("Year 1 combined DSCR", lambda e: number(e.metrics.year_one_dscr), False),
        ("Lowest DSCR over the hold", lambda e: number(e.metrics.min_dscr), False),
        ("Debt yield (NOI / debt)", lambda e: percent(e.metrics.debt_yield), False),
        ("Exit", None, False),
        ("Exit cap rate", lambda e: percent(e.deal.exit_cap_rate), False),
        ("Exit value", lambda e: money(e.metrics.exit_value), True),
        ("Less: selling costs", lambda e: money(e.metrics.selling_costs), False),
        (
            "Less: debt, prepayment & preferred",
            lambda e: money(
                e.metrics.loan_balance_at_exit
                + e.metrics.prepayment_penalties
                + e.metrics.preferred_equity_repayment
            ),
            False,
        ),
        ("Net sale proceeds", lambda e: money(e.metrics.net_sale_proceeds), True),
    ]
    _comparison_table(slide, entries, lines, size=10.5, basis_label=data.basis_label)


def _comparison_table(slide, entries, lines, size: float, basis_label: str | None = None) -> None:
    header = ["Deal"] + [f"{short_name(e.deal.name)}\n{_standing(e)}" for e in entries]
    rows = [header]
    section, bold, highlight = set(), set(), set()
    for label, extractor, is_bold in lines:
        index = len(rows)
        if extractor is None:
            section.add(index)
            rows.append([label.upper()] + [""] * len(entries))
            continue
        if is_bold:
            bold.add(index)
        if label == basis_label:
            highlight.add(index)
        rows.append([label] + [extractor(e) for e in entries])
    label_w = 3.0
    col_w = (CONTENT_W - label_w) / len(entries)
    available = 6.85 - CONTENT_TOP - 0.6
    row_h = min(0.36, available / (len(rows) - 1))
    _table(
        slide,
        MARGIN,
        CONTENT_TOP,
        [label_w] + [col_w] * len(entries),
        rows,
        row_h,
        size,
        section_rows=frozenset(section),
        bold_rows=frozenset(bold),
        highlight_rows=frozenset(highlight),
        header_h=0.55,
    )


def _add_hurdle_slide(deck: _Deck) -> None:
    data = deck.data
    slide = deck.content_slide(
        "Ranking basis",
        "How each hurdle was built: Treasury plus survey cap-rate spread plus NOI growth",
    )
    rows = [data.hurdle_headers] + data.hurdle_rows
    widths = [3.0, 1.7, 1.1, 1.2, 1.2, 1.1, 3.03]
    row_h = min(0.5, 4.2 / len(rows))
    _table(
        slide,
        MARGIN,
        CONTENT_TOP,
        widths,
        rows,
        row_h,
        12,
        aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * 4 + [PP_ALIGN.LEFT],
        bold_cols=frozenset({5}),
    )
    bottom = CONTENT_TOP + row_h * len(rows)
    _paragraphs(
        slide, MARGIN, bottom + 0.3, CONTENT_W, 1.2, [data.sources_note], size=11, color=MUTED
    )


def _add_profile_slide(deck: _Deck, entry: RankedEntry, index: int, total: int) -> None:
    deal = entry.deal
    facts = [
        _standing(entry) if entry.rank else "Excluded from ranking",
        format_assumption_value("sub_asset_class", deal.sub_asset_class),
        f"{deal.unit_count} units" if deal.unit_count is not None else f"{deal.bed_count} beds",
        f"{deal.hold_period_years}-year hold",
    ]
    slide = deck.content_slide(f"Deal profile {index} of {total}", f"Assumptions — {deal.name}")
    _text(slide, MARGIN, 1.32, CONTENT_W, 0.25, "  |  ".join(facts), 12, False, MUTED)
    _kpi_tiles(slide, 1.75, _kpis_for(entry), height=0.95)

    gap = 0.3
    col_w = (CONTENT_W - 2 * gap) / 3
    groups = dict(present_rows(deal, ASSUMPTION_GROUPS))
    columns = [
        ["Property & revenue"],
        ["Operating expenses", "Hold & exit"],
        ["Acquisition & financing"],
    ]
    for c, titles in enumerate(columns):
        rows: list[list[str]] = [["", ""]]
        sections: set[int] = set()
        for title in titles:
            if title not in groups:
                continue
            sections.add(len(rows))
            rows.append([title.upper(), ""])
            rows.extend([label, value] for label, value in groups[title])
        rows = rows[1:]  # no dark header: the section rows are the headers
        sections = {s - 1 for s in sections}
        row_h = min(0.27, 4.0 / max(len(rows), 1))
        _assumption_table(slide, MARGIN + c * (col_w + gap), 2.95, col_w, rows, row_h, sections)


def _assumption_table(slide, x, y, width, rows, row_h, sections) -> None:
    shape = slide.shapes.add_table(
        len(rows), 2, Inches(x), Inches(y), Inches(width), Inches(row_h * len(rows))
    )
    table = shape.table
    table.horz_banding = False
    table.first_row = False
    style_id = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style_id is not None:
        style_id.text = _NO_STYLE_NO_GRID
    table.columns[0].width = Inches(width * 0.62)
    table.columns[1].width = Inches(width * 0.38)
    for r, (label, value) in enumerate(rows):
        table.rows[r].height = Inches(row_h)
        if r in sections:
            table.cell(r, 0).merge(table.cell(r, 1))
            for c in range(2):
                _style_cell(
                    table.cell(r, c), label if c == 0 else "", 9, True, PRIMARY, TINT, border=None
                )
        else:
            _style_cell(table.cell(r, 0), label, 10.5)
            _style_cell(table.cell(r, 1), value, 10.5, True, align=PP_ALIGN.RIGHT)


def _add_capital_stack_slide(deck: _Deck, entry: RankedEntry) -> None:
    slide = deck.content_slide("Capital structure", f"Capital stack — {entry.deal.name}")
    segments = capital_stack_segments(entry.deal, entry.metrics)
    total = sum((amount for _, _, amount in segments), Decimal(0))

    bar_y, bar_h = 1.6, 0.55
    if total > 0:
        x = MARGIN
        for _name, group, amount in segments:
            width = CONTENT_W * float(amount / total)
            cell = _rect(slide, x, bar_y, max(width, 0.02), bar_h, _GROUP_COLORS[group])
            cell.line.color.rgb = WHITE
            cell.line.width = Pt(1.5)
            if width >= 0.9:
                tf = cell.text_frame
                tf.margin_left = tf.margin_right = Inches(0.04)
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                paragraph = tf.paragraphs[0]
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = percent(amount / total, 0)
                _font(run, 12, True, INK if group is LayerGroup.EQUITY else WHITE)
            x += width
        legend_x = MARGIN
        for label, group in (
            ("Senior debt", LayerGroup.SENIOR),
            ("Subordinate", LayerGroup.SUBORDINATE),
            ("Equity", LayerGroup.EQUITY),
        ):
            _rect(slide, legend_x, bar_y + bar_h + 0.14, 0.14, 0.14, _GROUP_COLORS[group])
            _text(slide, legend_x + 0.22, bar_y + bar_h + 0.1, 1.5, 0.25, label, 10, False, MUTED)
            legend_x += 1.7

    body = capital_stack_rows(entry.deal, entry.metrics)
    rows = [CAPITAL_STACK_HEADERS] + body
    top = 2.75
    row_h = min(0.42, (6.85 - top - 0.4) / (len(rows) - 1))
    _table(
        slide,
        MARGIN,
        top,
        [3.0, 1.5, 1.6, 1.1, 1.8, 3.33],
        rows,
        row_h,
        11 if len(rows) < 12 else 9,
        aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT]
        + [PP_ALIGN.LEFT, PP_ALIGN.LEFT],
        bold_rows=frozenset({len(rows) - 1}),
        header_h=0.4,
    )
    # Leverage read-outs under the table, when the stack is short enough to leave room.
    table_bottom = top + 0.4 + row_h * (len(rows) - 1)
    if table_bottom <= 5.3:
        m, deal = entry.metrics, entry.deal
        leverage = m.total_debt / deal.purchase_price if deal.purchase_price else None
        _kpi_tiles(
            slide,
            max(table_bottom + 0.45, 5.35),
            [
                ("Total debt", money(m.total_debt)),
                ("Debt / price", percent(leverage, 1)),
                ("Debt yield", percent(m.debt_yield)),
                ("Year 1 senior DSCR", number(m.year_one_senior_dscr)),
                ("Lowest DSCR", number(m.min_dscr)),
                ("Equity invested", money(m.equity_invested)),
            ],
            height=0.95,
        )


_MAX_PROFORMA_YEARS = 10


def _add_proforma_slide(deck: _Deck, entry: RankedEntry) -> None:
    flows = entry.metrics.annual_cash_flows[:_MAX_PROFORMA_YEARS]
    slide = deck.content_slide("Pro forma", f"Pro forma cash flow — {entry.deal.name}")
    _flow_chart(slide, entry, MARGIN, CONTENT_TOP, CONTENT_W, 2.7)

    lines = [
        ("Effective gross income", lambda f: f.egi, False),
        ("Operating expenses", lambda f: f.egi - f.noi, False),
        ("Net operating income", lambda f: f.noi, True),
        ("Replacement reserves", lambda f: f.replacement_reserves, False),
        ("Debt service", lambda f: f.debt_service, False),
        ("Cash flow after debt service", lambda f: f.levered_cash_flow, True),
    ]
    rows = [["$"] + [f"Yr {f.year}" for f in flows]]
    bold = set()
    for i, (label, getter, is_bold) in enumerate(lines, start=1):
        if is_bold:
            bold.add(i)
        rows.append([label] + [money(getter(f)) for f in flows])
    label_w = 2.4
    _table(
        slide,
        MARGIN,
        4.4,
        [label_w] + [(CONTENT_W - label_w) / len(flows)] * len(flows),
        rows,
        0.36,
        10 if len(flows) <= 7 else 9,
        bold_rows=frozenset(bold),
    )
    if len(entry.metrics.annual_cash_flows) > _MAX_PROFORMA_YEARS:
        _text(
            slide,
            MARGIN,
            4.4 + 0.36 * len(rows) + 0.1,
            CONTENT_W,
            0.25,
            f"Years 1–{_MAX_PROFORMA_YEARS} of {entry.deal.hold_period_years} shown.",
            10,
            False,
            MUTED,
            italic=True,
        )


def _add_methodology_slide(deck: _Deck) -> None:
    data = deck.data
    slide = deck.content_slide("Appendix", "Methodology and disclaimers")
    _text(slide, MARGIN, CONTENT_TOP, 6, 0.3, "Ranking basis", 14, True, PRIMARY)
    _paragraphs(
        slide,
        MARGIN,
        CONTENT_TOP + 0.4,
        CONTENT_W,
        1.8,
        [
            data.basis_note,
            "Both IRRs are shown for every deal. Ranking uses the unlevered IRR only; levered "
            "IRR is for reference and is never sorted on, so leverage cannot lift a deal in "
            "the ranking.",
        ],
        size=12,
        gap=8,
    )
    _text(slide, MARGIN, 3.9, 6, 0.3, "Disclaimers", 14, True, PRIMARY)
    _paragraphs(slide, MARGIN, 4.3, CONTENT_W, 2.4, [theme.DISCLAIMER], size=11, color=MUTED, gap=6)


def build_pptx_bytes(data: ReportData) -> bytes:
    deck = _Deck(data)
    _add_cover(deck)
    _add_summary_slide(deck)
    _add_ranking_slide(deck)
    _add_operating_slide(deck)
    _add_returns_slide(deck)
    if data.hurdle_rows:
        _add_hurdle_slide(deck)
    for index, entry in enumerate(data.entries, start=1):
        _add_profile_slide(deck, entry, index, len(data.entries))
        _add_capital_stack_slide(deck, entry)
        _add_proforma_slide(deck, entry)
    _add_methodology_slide(deck)

    buffer = io.BytesIO()
    deck.prs.save(buffer)
    return buffer.getvalue()
