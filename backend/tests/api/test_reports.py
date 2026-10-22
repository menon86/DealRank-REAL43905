"""See docs/build-plan.md Phase 6 — PDF and PPTX export. Every prior
API/DB code path this session had a real bug the first time it actually
ran (a masked password, an enum stored by name instead of value, an
async-only test transport); these exist to catch the equivalent for
export rather than trusting it on lint alone.
"""

from tests.api.payloads import STUDENT_HOUSING_PAYLOAD, SUBURBAN_GARDEN_PAYLOAD


def _seed_two_deals(client) -> list[str]:
    student = client.post("/deals", json=STUDENT_HOUSING_PAYLOAD).json()
    suburban = client.post("/deals", json=SUBURBAN_GARDEN_PAYLOAD).json()
    return [student["id"], suburban["id"]]


def test_pdf_report_returns_a_real_pdf(client):
    deal_ids = _seed_two_deals(client)

    resp = client.get(
        "/reports/pdf", params=[("deal_ids", i) for i in deal_ids] + [("hurdle_rate", "0.08")]
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF-")
    assert len(resp.content) > 500


def test_pptx_report_returns_a_real_pptx(client):
    deal_ids = _seed_two_deals(client)

    resp = client.get(
        "/reports/pptx", params=[("deal_ids", i) for i in deal_ids] + [("hurdle_rate", "0.08")]
    )
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    # PPTX/DOCX/XLSX files are zip archives — this is the zip magic number.
    assert resp.content.startswith(b"PK")
    assert len(resp.content) > 500


def test_pdf_report_unknown_deal_id_is_404(client):
    import uuid

    resp = client.get("/reports/pdf", params={"deal_ids": str(uuid.uuid4()), "hurdle_rate": "0.08"})
    assert resp.status_code == 404


def test_report_requires_at_least_one_deal_id(client):
    resp = client.get("/reports/pdf", params={"hurdle_rate": "0.08"})
    assert resp.status_code == 422


def test_pptx_tables_fit_on_their_slides_at_five_deals(client):
    """Every table's stored geometry has to end inside the slide. The
    assumptions table used to end at 7.65" on a 7.5" slide, which nothing
    caught because the file still opened; five deals is the most the UI
    lets you compare, so the comparison slide is checked at its widest."""
    import io

    from pptx import Presentation

    deal_ids = _seed_two_deals(client) + _seed_two_deals(client)
    deal_ids.append(client.post("/deals", json=STUDENT_HOUSING_PAYLOAD).json()["id"])

    resp = client.get(
        "/reports/pptx", params=[("deal_ids", i) for i in deal_ids] + [("hurdle_rate", "0.08")]
    )
    assert resp.status_code == 200

    prs = Presentation(io.BytesIO(resp.content))
    tables = [shape for slide in prs.slides for shape in slide.shapes if shape.has_table]
    assert tables
    for shape in tables:
        table = shape.table
        assert shape.top + sum(row.height for row in table.rows) <= prs.slide_height
        assert shape.left + sum(col.width for col in table.columns) <= prs.slide_width


def test_pdf_report_survives_markup_characters_in_deal_names(client):
    names = ["Smith & Sons Apts", "Bldg <A> Plaza", "<b>Bold</b> Place", '<img src="x"/>']
    ids = [
        client.post("/deals", json={**SUBURBAN_GARDEN_PAYLOAD, "name": name}).json()["id"]
        for name in names
    ]
    resp = client.get(
        "/reports/pdf", params=[("deal_ids", i) for i in ids] + [("hurdle_rate", "0.08")]
    )
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF-")


def _pdf_page_texts(pdf: bytes) -> list[str]:
    """Decoded content stream of each page, in page order. ReportLab writes
    one ASCII85 + Flate encoded content stream per page, with text drawn as
    `(string) Tj` operators, so a substring check per page is enough to tell
    which page a heading or table row landed on."""
    import base64
    import re
    import zlib

    streams = re.findall(
        rb"/Filter \[ /ASCII85Decode /FlateDecode \].*?stream\r?\n(.*?)~>\s*endstream",
        pdf,
        re.S,
    )
    return [zlib.decompress(base64.a85decode(s)).decode("latin-1") for s in streams]


def test_pdf_keeps_headings_with_tables_and_never_splits_a_table(client):
    """Five deals is the largest report the UI can ask for. Every heading
    must share a page with the start of its table, and each table must
    start and end on the same page."""
    deal_ids = _seed_two_deals(client) + _seed_two_deals(client)
    deal_ids.append(client.post("/deals", json=STUDENT_HOUSING_PAYLOAD).json()["id"])

    resp = client.get(
        "/reports/pdf", params=[("deal_ids", i) for i in deal_ids] + [("hurdle_rate", "0.08")]
    )
    assert resp.status_code == 200
    pages = _pdf_page_texts(resp.content)
    assert len(pages) >= 2

    def page_of(text: str) -> int:
        hits = [i for i, page in enumerate(pages) if f"({text})" in page]
        assert len(hits) == 1, f"{text!r} found on pages {hits}"
        return hits[0]

    # Comparison: heading, first row and last row on one page.
    assert page_of("Comparison") == page_of("Hold period") == page_of("Net sale proceeds")
    # Ranking: heading, header row and every ranked row on one page.
    assert page_of("Ranking") == page_of("Rank") == page_of("#1") == page_of("#5")
    # Assumptions heading sits with the first deal's table.
    assumptions_page = page_of("Assumptions")
    assert "(Sub-asset class)" in pages[assumptions_page]
    # Each assumptions table starts and ends on the same page.
    for page in pages:
        assert page.count("(Sub-asset class)") == page.count("(Selling costs rate)")


# ------------------------------------------------- risk-adjusted reports (the default)


def _pptx_text(content: bytes) -> str:
    import io

    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _risk_params(ids, **extra):
    return [("deal_ids", i) for i in ids] + [("treasury_yield", "0.046")] + list(extra.items())


def test_default_reports_use_the_risk_adjusted_ranking(client, spreads):
    from tests.api.payloads import URBAN_MIDRISE_PAYLOAD

    ids = _seed_two_deals(client) + [client.post("/deals", json=URBAN_MIDRISE_PAYLOAD).json()["id"]]
    ranking = client.post(
        "/rank/risk-adjusted", json={"deal_ids": ids, "treasury_yield": "0.046"}
    ).json()

    pdf = client.get("/reports/pdf", params=_risk_params(ids))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")

    resp = client.get("/reports/pptx", params=_risk_params(ids))
    assert resp.status_code == 200
    text = _pptx_text(resp.content)
    # The same order and the same numbers as the ranking endpoint.
    for entry in ranking["ranked"]:
        assert entry["deal"]["name"] in text
    first, second = ranking["ranked"][0]["deal"]["name"], ranking["ranked"][1]["deal"]["name"]
    assert text.index(first) < text.index(second)
    # The slide that says where the hurdle came from.
    assert "How each hurdle was built" in text
    assert "test survey" in text
    assert "override" in text
    # Excluded deals are named with the reason, not silently dropped.
    assert any(e["deal"]["name"] in text and e["reason"][:20] in text for e in ranking["excluded"])


def test_report_gate_parameters_change_the_report(client, spreads):
    ids = _seed_two_deals(client)
    lax = _pptx_text(
        client.get("/reports/pptx", params=_risk_params(ids, dscr_minimum="0")).content
    )
    strict = _pptx_text(
        client.get("/reports/pptx", params=_risk_params(ids, dscr_minimum="3")).content
    )
    assert "Excluded:" not in lax
    assert "Excluded:" in strict


def test_report_preference_weights_add_a_score_column(client, spreads):
    ids = _seed_two_deals(client)
    plain = _pptx_text(client.get("/reports/pptx", params=_risk_params(ids)).content)
    weighted = _pptx_text(
        client.get(
            "/reports/pptx",
            params=_risk_params(ids, weight_irr="1", weight_cash_flow="1", weight_risk="1"),
        ).content
    )
    assert "Score (0–100)" not in plain
    assert "Score (0–100)" in weighted


def test_report_without_reference_data_says_so(client):
    ids = _seed_two_deals(client)
    resp = client.get("/reports/pdf", params=_risk_params(ids))
    assert resp.status_code == 409


def test_flat_hurdle_report_still_available(client):
    ids = _seed_two_deals(client)
    resp = client.get(
        "/reports/pptx", params=[("deal_ids", i) for i in ids] + [("hurdle_rate", "0.08")]
    )
    text = _pptx_text(resp.content)
    assert "Hurdle rate: 8.00%" in text
    assert "How each hurdle was built" not in text


def test_new_underwriting_lines_appear_in_the_reports(client, spreads):
    payload = {
        **SUBURBAN_GARDEN_PAYLOAD,
        "management_fee_rate": "0.03",
        "replacement_reserve_per_unit_or_bed": "300",
    }
    ids = [client.post("/deals", json=payload).json()["id"]]
    text = _pptx_text(
        client.get("/reports/pptx", params=_risk_params(ids, dscr_minimum="0")).content
    )
    for label in (
        "Less: management fee",
        "Less: replacement reserves",
        "Lowest DSCR over the hold",
    ):
        assert label in text


def test_pptx_tables_fit_in_risk_mode_at_five_deals(client, spreads):
    import io

    from pptx import Presentation

    ids = _seed_two_deals(client) + _seed_two_deals(client)
    ids.append(client.post("/deals", json=STUDENT_HOUSING_PAYLOAD).json()["id"])
    resp = client.get("/reports/pptx", params=_risk_params(ids, dscr_minimum="0"))
    prs = Presentation(io.BytesIO(resp.content))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                assert shape.top + sum(r.height for r in table.rows) <= prs.slide_height
                assert shape.left + sum(c.width for c in table.columns) <= prs.slide_width
