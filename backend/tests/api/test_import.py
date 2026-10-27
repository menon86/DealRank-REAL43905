"""GET /deals/import/template and POST /deals/import."""

import io

from openpyxl import load_workbook


def test_template_downloads_as_xlsx(client):
    resp = client.get("/deals/import/template")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "dealrank_underwriting_template.xlsx" in resp.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Instructions", "Property", "Unit Mix", "Financing", "Lists"]
    assert workbook["Lists"].sheet_state == "hidden"


def test_template_imports_back(client):
    template = client.get("/deals/import/template").content
    resp = client.post(
        "/deals/import",
        files={"file": ("dealrank_underwriting_template.xlsx", template)},
    )
    assert resp.status_code == 200
    body = resp.json()
    by_key = {f["key"]: f for f in body["fields"]}
    assert by_key["purchase_price"]["value"] == "24500000"
    assert by_key["purchase_price"]["source"] == {
        "sheet": "Property",
        "cell": "B9",
        "label": "Purchase price",
    }
    assert by_key["exit_cap_rate_pct"]["value"] == "5.75"
    assert body["unit_mix"]["sheet"] == "Unit Mix"


def test_csv_upload(client):
    resp = client.post(
        "/deals/import",
        files={"file": ("deal.csv", b'Purchase Price,"$1,000,000"\n')},
    )
    assert resp.status_code == 200
    by_key = {f["key"]: f for f in resp.json()["fields"]}
    assert by_key["purchase_price"]["value"] == "1000000"


def test_unreadable_upload_is_a_422_with_a_message(client):
    resp = client.post("/deals/import", files={"file": ("deal.xls", b"\xd0\xcf\x11\xe0")})
    assert resp.status_code == 422
    assert "save it as .xlsx" in resp.json()["detail"]


def test_import_routes_dont_shadow_deal_routes(client):
    assert client.get("/deals").status_code == 200
