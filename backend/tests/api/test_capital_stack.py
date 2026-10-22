"""Capital stack through the API: layers persist and round-trip, the
single-loan fields follow the senior layer, invalid stacks are a 422, and
the metrics/report endpoints reflect every layer.
"""

import io
from decimal import Decimal

from tests.api.payloads import SUBURBAN_GARDEN_PAYLOAD

STACK = [
    {
        "type": "bridge_debt_fund",
        "amount": "18000000",
        "rate_type": "floating",
        "index_name": "SOFR",
        "index_rate": "0.043",
        "spread": "0.0275",
        "io_months": 24,
        "amortization_years": 30,
        "term_years": 5,
        "recourse": False,
        "rate_cap_cost": "120000",
        "origination_fee_rate": "0.01",
    },
    {"type": "mezzanine", "amount": "2000000", "interest_rate": "0.11"},
    {"type": "preferred_equity", "amount": "1500000", "preferred_return": "0.09"},
    {"type": "lp_syndication", "amount": "6000000"},
    {"type": "other_equity", "label": "Family office co-invest", "amount": "2196000"},
]


def _create(client, **overrides):
    return client.post("/deals", json={**SUBURBAN_GARDEN_PAYLOAD, **overrides})


def test_stack_round_trips_and_drives_the_single_loan_fields(client):
    resp = _create(client, capital_stack=STACK)
    assert resp.status_code == 201, resp.text
    deal = resp.json()

    # The senior layer is mirrored onto loan_amount / interest_rate /
    # amortization_years, whatever the payload sent for them.
    assert Decimal(deal["loan_amount"]) == Decimal("18000000")
    assert Decimal(deal["interest_rate"]) == Decimal("0.0705")
    assert deal["amortization_years"] == 30

    fetched = client.get(f"/deals/{deal['id']}").json()
    assert [layer["type"] for layer in fetched["capital_stack"]] == [
        layer["type"] for layer in STACK
    ]
    assert fetched["capital_stack"][4]["label"] == "Family office co-invest"
    assert fetched["capital_stack"][0]["recourse"] is False


def test_metrics_reflect_every_layer(client):
    plain_id = _create(client).json()["id"]
    stacked_id = _create(client, capital_stack=STACK).json()["id"]
    plain = client.get(f"/deals/{plain_id}/metrics").json()
    stacked = client.get(f"/deals/{stacked_id}/metrics").json()

    assert Decimal(stacked["total_debt"]) == Decimal("20000000")
    assert Decimal(stacked["preferred_equity"]) == Decimal("1500000")
    assert Decimal(stacked["upfront_financing_costs"]) == Decimal("300000")
    assert Decimal(stacked["year_one_dscr"]) < Decimal(stacked["year_one_senior_dscr"])
    year_one = stacked["annual_cash_flows"][0]
    assert Decimal(year_one["preferred_distribution"]) == Decimal("135000")
    # Five-year bridge term inside a ten-year hold: balloon in year 5.
    assert Decimal(stacked["annual_cash_flows"][4]["balloon_payment"]) > 0

    # Financing never moves the ranking basis.
    assert stacked["unlevered_irr"] == plain["unlevered_irr"]


def test_invalid_stack_is_a_422(client):
    two_seniors = [STACK[0], {**STACK[0], "type": "cmbs"}, STACK[3]]
    resp = _create(client, capital_stack=two_seniors)
    assert resp.status_code == 422
    assert "Only one senior debt layer" in resp.text

    no_equity = [STACK[0]]
    assert _create(client, capital_stack=no_equity).status_code == 422

    unnamed_other = [STACK[0], {"type": "other_equity", "amount": "1"}]
    assert _create(client, capital_stack=unnamed_other).status_code == 422


def test_all_cash_stack_zeroes_the_loan(client):
    cash = [{"type": "all_cash", "amount": "29376000"}]
    deal = _create(client, capital_stack=cash).json()
    assert Decimal(deal["loan_amount"]) == 0
    metrics = client.get(f"/deals/{deal['id']}/metrics").json()
    assert metrics["year_one_dscr"] is None
    assert metrics["year_one_senior_dscr"] is None
    assert metrics["levered_irr"] == metrics["unlevered_irr"]


def test_patch_adds_and_clears_a_stack(client):
    deal_id = _create(client).json()["id"]

    senior = {"type": "life_company", "amount": "15000000", "interest_rate": "0.055"}
    equity = {"type": "sponsor_gp", "amount": "0"}
    resp = client.patch(f"/deals/{deal_id}", json={"capital_stack": [senior, equity]})
    assert resp.status_code == 200, resp.text
    patched = resp.json()
    assert Decimal(patched["loan_amount"]) == Decimal("15000000")
    assert Decimal(patched["interest_rate"]) == Decimal("0.055")

    resp = client.patch(f"/deals/{deal_id}", json={"capital_stack": [senior]})
    assert resp.status_code == 422

    resp = client.patch(f"/deals/{deal_id}", json={"capital_stack": None})
    assert resp.status_code == 200
    assert resp.json()["capital_stack"] is None


def test_reports_show_the_capital_stack_with_other_labels_as_typed(client):
    from pptx import Presentation

    ids = [
        _create(client, capital_stack=STACK).json()["id"],
        _create(client, name="Plain loan deal").json()["id"],
    ]
    params = [("deal_ids", i) for i in ids] + [("hurdle_rate", "0.08")]

    assert client.get("/reports/pdf", params=params).status_code == 200

    resp = client.get("/reports/pptx", params=params)
    assert resp.status_code == 200
    prs = Presentation(io.BytesIO(resp.content))
    cells = [
        cell.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    ]
    assert "Family office co-invest" in cells
    assert "Bridge / debt fund" in cells
    assert "Senior loan (type not specified)" in cells
    assert "Year 1 senior DSCR" in cells
    assert "Year 1 combined DSCR" in cells


# ------------------------------------------------- common equity must exist


def test_stack_whose_debt_covers_the_cost_is_rejected(client):
    stack = [
        {
            "type": "agency_fannie_mae",
            "amount": "20000000",
            "interest_rate": "0.06",
            "amortization_years": 30,
        },
        {"type": "mezzanine", "amount": "9000000", "interest_rate": "0.12"},
        {
            "type": "supplemental_loan",
            "amount": "4000000",
            "interest_rate": "0.08",
            "amortization_years": 25,
        },
        {"type": "sponsor_gp", "amount": "1"},
    ]
    payload = {**SUBURBAN_GARDEN_PAYLOAD, "purchase_price": "22000000", "capital_stack": stack}
    resp = client.post("/deals", json=payload)
    assert resp.status_code == 422
    assert "common equity" in resp.text


def test_preferred_equity_counts_against_common_equity(client):
    stack = [
        {
            "type": "agency_fannie_mae",
            "amount": "18000000",
            "interest_rate": "0.06",
            "amortization_years": 30,
        },
        {"type": "preferred_equity", "amount": "10500000", "preferred_return": "0.09"},
        {"type": "sponsor_gp", "amount": "1"},
    ]
    resp = client.post("/deals", json={**SUBURBAN_GARDEN_PAYLOAD, "capital_stack": stack})
    assert resp.status_code == 422


def test_single_loan_leaving_under_five_percent_equity_is_rejected(client):
    payload = {
        **SUBURBAN_GARDEN_PAYLOAD,
        "purchase_price": "10000000",
        "closing_costs": "100000",
        "loan_amount": "10000000",
    }
    assert client.post("/deals", json=payload).status_code == 422


def test_patch_cannot_push_a_deal_below_the_equity_floor(client):
    deal = client.post("/deals", json=SUBURBAN_GARDEN_PAYLOAD).json()
    resp = client.patch(f"/deals/{deal['id']}", json={"loan_amount": "28800000"})
    assert resp.status_code == 422


def test_rates_keep_sub_basis_point_precision(client):
    payload = {
        **SUBURBAN_GARDEN_PAYLOAD,
        "interest_rate": "0.06125",
        "exit_cap_rate": "0.05125",
    }
    created = client.post("/deals", json=payload).json()
    fetched = client.get(f"/deals/{created['id']}").json()
    assert Decimal(fetched["interest_rate"]) == Decimal("0.061250")
    assert Decimal(fetched["exit_cap_rate"]) == Decimal("0.051250")


HUD_STACK = [
    {
        "type": "hud_223f",
        "amount": "18000000",
        "interest_rate": "0.0575",
        "amortization_years": 35,
        "term_years": 35,
        "mip_annual_rate": "0.0025",
        "mip_upfront_rate": "0.0025",
        "prepayment_type": "step_down",
        "prepayment_schedule": ["0.10", "0.09", "0.08", "0.07", "0.06", "0.05", "0.04"],
    },
    {"type": "sponsor_gp", "amount": "11376000"},
]


def test_hud_insurance_and_prepayment_round_trip_and_reach_the_metrics(client):
    resp = _create(client, capital_stack=HUD_STACK)
    assert resp.status_code == 201, resp.text
    layer = client.get(f"/deals/{resp.json()['id']}").json()["capital_stack"][0]
    assert layer["prepayment_type"] == "step_down"
    assert [Decimal(r) for r in layer["prepayment_schedule"]][0] == Decimal("0.10")
    assert Decimal(layer["mip_annual_rate"]) == Decimal("0.0025")

    metrics = client.get(f"/deals/{resp.json()['id']}/metrics").json()
    assert Decimal(metrics["annual_cash_flows"][0]["mortgage_insurance"]) > 0
    # Ten-year hold, seven-year schedule: no penalty left by the sale.
    assert Decimal(metrics["prepayment_penalties"]) == 0
    assert Decimal(metrics["upfront_financing_costs"]) == Decimal("45000")


def test_yield_maintenance_penalty_reduces_sale_proceeds(client):
    agency = {
        "type": "agency_fannie_mae",
        "amount": "18000000",
        "interest_rate": "0.06",
        "amortization_years": 30,
        "term_years": 12,
    }
    plain_id = _create(client, capital_stack=[agency, STACK[3]]).json()["id"]
    ym = {**agency, "prepayment_type": "yield_maintenance", "prepayment_treasury_rate": "0.04"}
    ym_id = _create(client, capital_stack=[ym, STACK[3]]).json()["id"]

    plain = client.get(f"/deals/{plain_id}/metrics").json()
    metrics = client.get(f"/deals/{ym_id}/metrics").json()
    penalty = Decimal(metrics["prepayment_penalties"])
    assert penalty > 0
    assert Decimal(metrics["net_sale_proceeds"]) == Decimal(plain["net_sale_proceeds"]) - penalty
    assert metrics["unlevered_irr"] == plain["unlevered_irr"]


def test_mortgage_insurance_on_a_non_hud_loan_is_a_422(client):
    agency = {**STACK[0], "mip_annual_rate": "0.0025"}
    resp = _create(client, capital_stack=[agency, STACK[3]])
    assert resp.status_code == 422
    assert "mortgage insurance applies only" in resp.text
